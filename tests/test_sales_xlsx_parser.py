from datetime import date
from io import BytesIO

from openpyxl import Workbook
import psycopg
from psycopg.rows import dict_row

from app.accounting import _save_weekly_report_expenses
from app.config import get_settings
from app.sales_xlsx import parse_sales_xlsx
from apoyo_db import conexion_pruebas


def make_sales_workbook() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "SEMANA 32"
    sheet.append([None, "EFECTIVO", "DOLARES", "TIPO DE CAMBIO", "CONVERSION", "PLATAFORMAS", "UBER", "VISA", "AMEX", "TARJETA DE CREDITO", "VENTA TOTAL"])
    dates = [date(2026, 8, day) for day in range(3, 10)]
    for index, current in enumerate(dates, start=1):
        sheet.append([current, 100 * index, 0, 0, 0, 10 * index, 20 * index, 30 * index, 40 * index, 70 * index, 200 * index])
    sheet.append([])
    sheet.append([None, *dates[:4]])
    for label, base in (("OTRO", 10), ("APLICACIONES", 20), ("ALIMENTOS", 120), ("BEBIDAS", 50)):
        sheet.append([label, *(base * index for index in range(1, 5))])
    sheet.append([None, *dates[4:]])
    for label, base in (("OTRO", 10), ("APLICACIONES", 20), ("ALIMENTOS", 120), ("BEBIDAS", 50)):
        sheet.append([label, *(base * index for index in range(5, 8))])
    sheet.cell(row=10, column=13, value="VISITAS")
    for row, visits in enumerate(range(100, 107), start=11):
        sheet.cell(row=row, column=13, value=visits)
    sheet.cell(row=30, column=1, value="GASTOS BARRA")
    for column, value in enumerate(("FECHA", "MONTO", "PAGO", "SEMANA", "No. FACTURA", "DESCRIPCION"), start=3):
        sheet.cell(row=30, column=column, value=value)
    sheet.cell(row=31, column=1, value="BOLA DE ORO")
    for column, value in enumerate((date(2026, 8, 5), 550, "EFECTIVO", 32, None, "1 KILO DE CAFÉ"), start=3):
        sheet.cell(row=31, column=column, value=value)
    sheet.cell(row=34, column=1, value="TOTAL")
    sheet.cell(row=40, column=1, value="GASTOS RESTAURANTE")
    for column, value in enumerate(("FECHA", "MONTO", "PAGO", "SEMANA", "No. FACTURA", "DESCRIPCION"), start=3):
        sheet.cell(row=40, column=column, value=value)
    sheet.cell(row=41, column=1, value="NORTEÑA")
    for column, value in enumerate((date(2026, 8, 6), 210, "EFECTIVO", 32, "A-9", "TORTILLA DE MAÍZ"), start=3):
        sheet.cell(row=41, column=column, value=value)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_xlsx_parser_accepts_plataformas_as_rappi_and_ignores_extra_cells():
    parsed = parse_sales_xlsx(make_sales_workbook(), expected_year=2026, expected_week=32)

    assert parsed["semana_iso"] == 32
    assert len(parsed["filas"]) == 7
    assert parsed["filas"][0]["rappi"] == 10
    assert parsed["filas"][0]["monto_neto"] == 200
    assert parsed["filas"][0]["venta_alimentos"] == 120
    assert parsed["filas"][6]["venta_bebidas"] == 350
    assert parsed["filas"][0]["visitas"] == 100
    assert parsed["proveedor_ocr"] == "xlsx-estructurado"
    assert [(item["area"], item["monto"]) for item in parsed["gastos"]] == [
        ("barra", 550), ("restaurante", 210),
    ]


def test_weekly_expense_import_is_idempotent():
    expense = {
        "fila_origen": 31, "area": "barra", "proveedor": "PYTEST CAFÉ SEMANAL",
        "fecha": date(2026, 8, 5), "monto": 550, "pago": "efectivo",
        "semana": 32, "numero_factura": None, "descripcion": "1 KILO DE CAFÉ",
        "incluir": True,
    }
    with conexion_pruebas(row_factory=dict_row) as conn:
        branch_id = conn.execute("SELECT id FROM catalog.sucursales WHERE clave = 'DEMO-A'").fetchone()["id"]
        first = _save_weekly_report_expenses(
            conn, branch_id=branch_id, expenses=[expense], actor="pytest",
            source_filename="semana-32.xlsx",
        )
        second = _save_weekly_report_expenses(
            conn, branch_id=branch_id, expenses=[expense], actor="pytest",
            source_filename="semana-32.xlsx",
        )
        assert first == {"created": 1, "existing": 0}
        assert second == {"created": 0, "existing": 1}
        conn.rollback()
