"""Plantilla XLSX del catálogo de proveedores, generada al vuelo."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

PROVIDER_HEADERS = [
    "rfc", "razon_social", "nombre_normalizado", "nombre_comercial", "categoria",
    "subcategoria", "aliases", "activo",
]
PROVIDER_ROWS = 200

FIELDS = [
    ("rfc", "No", "RFC del emisor; dejar vacío para proveedor informal."),
    ("razon_social", "Sí", "Razón social o nombre legal exacto."),
    ("nombre_normalizado", "Sí", "Nombre único que verá el capturista y Analytics."),
    ("nombre_comercial", "No", "Nombre de fachada o marca."),
    ("categoria", "No", "Clave de la hoja Categorias. Nómina y Comisiones no se capturan con facturas."),
    ("subcategoria", "No", "Clave de la hoja Subcategorias; debe corresponder a la categoría."),
    ("aliases", "No", "Variantes OCR o nombres históricos separados por |."),
    ("activo", "Sí", "SI o NO."),
]

HEADER_FILL = PatternFill("solid", fgColor="174C36")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _sheet(workbook: Workbook, title: str, headers: list[str], rows: list[tuple], widths: list[int]):
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for row in rows:
        sheet.append(list(row))
    for cell in sheet[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for index, width in enumerate(widths):
        sheet.column_dimensions[chr(65 + index)].width = width
    sheet.freeze_panes = "A2"
    return sheet


def build_provider_template(
    categories: list[tuple[str, str]], subcategories: list[tuple[str, str, str]]
) -> bytes:
    workbook = Workbook()
    guide = workbook.active
    guide.title = "Instrucciones"
    guide["A1"] = "Plantilla maestra · Proveedores"
    guide["A1"].font = Font(bold=True, size=16, color="174C36")
    steps = [
        "Llene una fila por proveedor en la hoja Proveedores.",
        "Separe alias con el carácter |.",
        "Use únicamente claves disponibles en las hojas de catálogos.",
        "Cargue el archivo en Catálogo y aprendizaje y revise la vista previa antes de confirmar.",
        "No cambie los encabezados de la hoja Proveedores: la app los usa para validar el archivo.",
    ]
    for offset, step in enumerate(steps, start=3):
        guide[f"A{offset}"] = f"{offset - 2}. {step}"
    guide["A9"], guide["B9"], guide["C9"] = "Campo", "Requerido", "Descripción"
    for cell in guide[9]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    for offset, field in enumerate(FIELDS, start=10):
        for column, value in zip("ABC", field):
            guide[f"{column}{offset}"] = value
    guide.column_dimensions["A"].width = 24
    guide.column_dimensions["B"].width = 12
    guide.column_dimensions["C"].width = 70

    providers = _sheet(
        workbook, "Proveedores", PROVIDER_HEADERS, [],
        [18, 30, 30, 30, 22, 22, 46, 10],
    )
    last = PROVIDER_ROWS + 1
    validations = [
        (f"=Categorias!$A$2:$A${len(categories) + 1}", f"E2:E{last}"),
        (f"=Subcategorias!$B$2:$B${len(subcategories) + 1}", f"F2:F{last}"),
        ('"SI,NO"', f"H2:H{last}"),
    ]
    for formula, cells in validations:
        validation = DataValidation(type="list", formula1=formula, allow_blank=True)
        providers.add_data_validation(validation)
        validation.add(cells)

    _sheet(workbook, "Sucursales", ["clave", "razon_social", "rfc", "restaurante", "nombre_sucursal", "aliases"], [], [16, 38, 18, 23, 23, 48])
    _sheet(workbook, "Categorias", ["clave", "nombre"], categories, [24, 28])
    _sheet(workbook, "Subcategorias", ["categoria", "clave", "nombre"], subcategories, [26, 24, 28])

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
