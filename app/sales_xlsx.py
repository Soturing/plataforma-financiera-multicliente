from __future__ import annotations

import io
import re
import unicodedata
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from app.errors import DomainError


MONTHS = {
    "ENE": 1, "FEB": 2, "MAR": 3, "ABR": 4, "MAY": 5, "JUN": 6,
    "JUL": 7, "AGO": 8, "SEP": 9, "OCT": 10, "NOV": 11, "DIC": 12,
}


def _plain(value: Any) -> str:
    decomposed = unicodedata.normalize("NFKD", str(value or "").upper())
    plain = "".join(char for char in decomposed if not unicodedata.combining(char))
    return " ".join(re.sub(r"[^A-Z0-9]+", " ", plain).split())


def _money(value: Any) -> Decimal:
    if value in (None, "", "-"):
        return Decimal("0")
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value)).quantize(Decimal("0.01"))
    cleaned = re.sub(r"[^0-9,.-]", "", str(value))
    if cleaned.count(",") and cleaned.count("."):
        cleaned = cleaned.replace(",", "")
    elif cleaned.count(",") == 1 and len(cleaned.rsplit(",", 1)[1]) == 2:
        cleaned = cleaned.replace(",", ".")
    else:
        cleaned = cleaned.replace(",", "")
    try:
        return Decimal(cleaned or "0").quantize(Decimal("0.01"))
    except InvalidOperation:
        return Decimal("0")


def _date(value: Any, year: int) -> date | None:
    if isinstance(value, (int, float)):
        try:
            converted = from_excel(value)
            return converted.date() if isinstance(converted, datetime) else converted
        except (TypeError, ValueError):
            return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw_text = str(value or "").strip()
    text = _plain(raw_text)
    match = re.search(r"\b(\d{1,2})\s+([A-Z]{3,})\b", text)
    if match and (month := MONTHS.get(match.group(2)[:3])):
        try:
            return date(year, month, int(match.group(1)))
        except ValueError:
            return None
    if match := re.match(r"^(\d{1,2})[/-](\d{1,2})[/-](20\d{2})$", raw_text):
        day, month, parsed_year = (int(item) for item in match.groups())
        try:
            return date(parsed_year, month, day)
        except ValueError:
            return None
    if match := re.match(r"^(20\d{2})-(\d{1,2})-(\d{1,2})$", raw_text):
        try:
            return date(*(int(item) for item in match.groups()))
        except ValueError:
            return None
    return None


def _find_header(matrix: list[list[Any]]) -> tuple[int, dict[str, int]]:
    aliases = {
        "efectivo": {"EFECTIVO"},
        "dolares": {"DOLARES"},
        "tipo_cambio": {"TIPO DE CAMBIO"},
        "conversion_usd": {"CONVERSION"},
        "rappi": {"RAPPI", "PLATAFORMAS"},
        "uber": {"UBER"},
        "visa": {"VISA"},
        "amex": {"AMEX"},
        "credito": {"CREDITO", "TARJETA DE CREDITO"},
        "monto_neto": {"VENTA TOTAL", "TOTAL VENTA"},
    }
    for row_index, row in enumerate(matrix[:50]):
        found: dict[str, int] = {}
        for column, value in enumerate(row):
            normalized = _plain(value)
            for field, names in aliases.items():
                if normalized in names:
                    found[field] = column
        if {"efectivo", "rappi", "uber", "visa", "amex", "monto_neto"}.issubset(found):
            return row_index, found
    raise DomainError("No se encontró el encabezado de ventas (EFECTIVO a VENTA TOTAL).", 422)


def parse_sales_xlsx(data: bytes, *, expected_year: int, expected_week: int | None = None) -> dict[str, Any]:
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise DomainError(f"No fue posible abrir el XLSX: {exc}", 422) from exc

    selected: tuple[list[list[Any]], int, dict[str, int]] | None = None
    for sheet in workbook.worksheets:
        matrix = [
            list(row)
            for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 500), values_only=True)
        ]
        try:
            header_row, headers = _find_header(matrix)
            selected = matrix, header_row, headers
            break
        except DomainError:
            continue
    if not selected:
        raise DomainError("Ninguna hoja contiene el formato semanal de ventas.", 422)
    matrix, header_row, headers = selected
    date_column = max(0, headers["efectivo"] - 1)
    rows: list[dict[str, Any]] = []
    for source in matrix[header_row + 1:header_row + 13]:
        parsed_date = _date(source[date_column] if date_column < len(source) else None, expected_year)
        if not parsed_date:
            continue
        row = {
            "fecha": parsed_date,
            "efectivo": Decimal("0"), "dolares": Decimal("0"),
            "tipo_cambio": Decimal("0"), "rappi": Decimal("0"), "uber": Decimal("0"),
            "visa": Decimal("0"), "amex": Decimal("0"), "monto_neto": Decimal("0"),
            "venta_otro": Decimal("0"), "venta_aplicaciones": Decimal("0"),
            "venta_alimentos": Decimal("0"), "venta_bebidas": Decimal("0"),
            "visitas": 0, "faltante": Decimal("0"), "sobrante": Decimal("0"),
            "total_capturado_manual": False,
        }
        for field in row:
            if field in headers and headers[field] < len(source):
                row[field] = _money(source[headers[field]])
        calculated = sum(
            (Decimal(row[field]) for field in ("efectivo", "rappi", "uber", "visa", "amex")),
            Decimal(row["dolares"]) * Decimal(row["tipo_cambio"]),
        )
        row["total_capturado_manual"] = abs(Decimal(row["monto_neto"]) - calculated) > Decimal("0.02")
        rows.append(row)
        if len(rows) == 7:
            break
    if len(rows) != 7:
        raise DomainError("El XLSX no contiene siete fechas válidas bajo el encabezado.", 422)

    by_date = {row["fecha"]: row for row in rows}
    current_dates: dict[int, date] = {}
    category_fields = {
        "OTRO": "venta_otro", "APLICACIONES": "venta_aplicaciones",
        "ALIMENTOS": "venta_alimentos", "BEBIDAS": "venta_bebidas",
    }
    for source in matrix[header_row + 8:]:
        dated_columns = {
            column: parsed for column, value in enumerate(source)
            if (parsed := _date(value, expected_year)) in by_date
        }
        if len(dated_columns) >= 2:
            current_dates = dated_columns
            continue
        label = next((_plain(value) for value in source if _plain(value) in category_fields), None)
        if label and current_dates:
            field = category_fields[label]
            for column, parsed_date in current_dates.items():
                if column < len(source):
                    by_date[parsed_date][field] = _money(source[column])

    for field, label in (("visitas", "VISITAS"), ("faltante", "FALTANTE"), ("sobrante", "SOBRANTE")):
        header_position = next(
            ((row_index, column) for row_index, source in enumerate(matrix) for column, value in enumerate(source) if _plain(value) == label),
            None,
        )
        if not header_position:
            continue
        source_row, column = header_position
        values = []
        for source in matrix[source_row + 1:source_row + 11]:
            value = source[column] if column < len(source) else None
            if value not in (None, ""):
                values.append(int(_money(value)) if field == "visitas" else _money(value))
            if len(values) == 7:
                break
        for row, value in zip(rows, values):
            row[field] = value

    inferred = rows[0]["fecha"].isocalendar()
    warnings = []
    if expected_week and inferred.week != expected_week:
        warnings.append(
            f"Las fechas del archivo corresponden a la semana {inferred.week}; la pantalla estaba en la {expected_week}."
        )
    payment_total = sum((Decimal(row["monto_neto"]) for row in rows), Decimal("0"))
    mix_total = sum(
        (sum((Decimal(row[field]) for field in category_fields.values()), Decimal("0")) for row in rows),
        Decimal("0"),
    )
    if mix_total and abs(payment_total - mix_total) > Decimal("0.02"):
        warnings.append(f"El desglose comercial difiere de la venta neta por {payment_total - mix_total:.2f}.")

    expenses: list[dict[str, Any]] = []
    area = None
    for row_number, source in enumerate(matrix, start=1):
        label = _plain(source[0] if source else None)
        if label == "GASTOS BARRA":
            area = "barra"
            continue
        if label == "GASTOS RESTAURANTE":
            area = "restaurante"
            continue
        if not area:
            continue
        if label == "TOTAL" or (label.startswith("GASTOS ") and label not in {"GASTOS BARRA", "GASTOS RESTAURANTE"}):
            area = None
            continue
        provider = str(source[0] or "").strip() if len(source) > 0 else ""
        expense_date = _date(source[2] if len(source) > 2 else None, expected_year)
        amount = _money(source[3] if len(source) > 3 else None)
        if not provider or not expense_date or amount <= 0:
            continue
        expenses.append({
            "fila_origen": row_number,
            "area": area,
            "proveedor": provider,
            "fecha": expense_date,
            "monto": amount,
            "pago": str(source[4] or "").strip().lower() if len(source) > 4 else "otro",
            "semana": int(_money(source[5])) if len(source) > 5 and source[5] not in (None, "") else expense_date.isocalendar().week,
            "numero_factura": str(source[6] or "").strip() or None if len(source) > 6 else None,
            "descripcion": str(source[7] or "").strip() or "Compra en efectivo" if len(source) > 7 else "Compra en efectivo",
            "incluir": True,
        })
    return {
        "anio_iso": inferred.year,
        "semana_iso": inferred.week,
        "filas": rows,
        "confianza_ocr": 1.0,
        "proveedor_ocr": "xlsx-estructurado",
        "advertencias": warnings,
        "gastos": expenses,
        "conciliacion": {
            "venta_neta": payment_total,
            "total_desglose": mix_total,
            "diferencia": mix_total - payment_total,
        },
    }
