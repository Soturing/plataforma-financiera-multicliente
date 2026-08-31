from __future__ import annotations

from pathlib import Path
from typing import Iterable

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "deliverables"
FIGURES_DIR = OUTPUT_DIR / "executive_brief_assets"
OUTPUT_PATH = OUTPUT_DIR / "propuesta-ejecutiva-automatizacion-financiera.docx"

# standard_business_brief with named override: internal restaurant green/gold palette.
GREEN = "176B45"
GREEN_DARK = "0F5234"
GREEN_SOFT = "E8F3ED"
GOLD = "D6A928"
GOLD_SOFT = "FFF3DC"
INK = "17211C"
MUTED = "68746D"
LINE = "DFE6E1"
LIGHT = "F3F5F2"
WHITE = "FFFFFF"
RED = "B33A3A"

PIL_FONT = Path("C:/Windows/Fonts/segoeui.ttf")
PIL_BOLD = Path("C:/Windows/Fonts/segoeuib.ttf")


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(str(PIL_BOLD if bold else PIL_FONT), size)


def wrap_lines(draw: ImageDraw.ImageDraw, text: str, selected_font, max_width: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    current = ""
    for word in words:
        candidate = f"{current} {word}".strip()
        if draw.textlength(candidate, font=selected_font) <= max_width or not current:
            current = candidate
        else:
            lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def centered_text(draw, box, text, selected_font, fill, max_width=None, spacing=8):
    x1, y1, x2, y2 = box
    width = max_width or (x2 - x1 - 30)
    lines = wrap_lines(draw, text, selected_font, width)
    bbox = draw.multiline_textbbox((0, 0), "\n".join(lines), font=selected_font, spacing=spacing, align="center")
    text_w = bbox[2] - bbox[0]
    text_h = bbox[3] - bbox[1]
    draw.multiline_text(
        ((x1 + x2 - text_w) / 2, (y1 + y2 - text_h) / 2),
        "\n".join(lines),
        font=selected_font,
        fill=fill,
        spacing=spacing,
        align="center",
    )


def node(draw, box, title, subtitle=None, fill="#FFFFFF", outline="#176B45", title_color="#17211C"):
    draw.rounded_rectangle(box, radius=22, fill=fill, outline=outline, width=4)
    x1, y1, x2, y2 = box
    if subtitle:
        centered_text(draw, (x1 + 20, y1 + 12, x2 - 20, y1 + (y2 - y1) * 0.58), title, font(29, True), title_color)
        centered_text(draw, (x1 + 25, y1 + (y2 - y1) * 0.52, x2 - 25, y2 - 8), subtitle, font(21), "#68746D")
    else:
        centered_text(draw, (x1 + 15, y1 + 10, x2 - 15, y2 - 10), title, font(27, True), title_color)


def arrow(draw, start, end, color="#176B45", width=7):
    draw.line([start, end], fill=color, width=width)
    x2, y2 = end
    x1, y1 = start
    dx, dy = x2 - x1, y2 - y1
    length = max((dx * dx + dy * dy) ** 0.5, 1)
    ux, uy = dx / length, dy / length
    px, py = -uy, ux
    size = 19
    p1 = (x2, y2)
    p2 = (x2 - ux * size + px * size * 0.55, y2 - uy * size + py * size * 0.55)
    p3 = (x2 - ux * size - px * size * 0.55, y2 - uy * size - py * size * 0.55)
    draw.polygon([p1, p2, p3], fill=color)


def base_canvas(width=1600, height=850):
    image = Image.new("RGB", (width, height), "#FFFFFF")
    return image, ImageDraw.Draw(image)


def save_manual_flow(path: Path):
    image, draw = base_canvas(1600, 760)
    draw.text((70, 45), "Flujo actual: doble captura y cierre dependiente de Excel", font=font(42, True), fill="#17211C")
    node(draw, (70, 180, 345, 330), "Facturas", "Documentos apilados")
    node(draw, (430, 180, 705, 330), "Captura manual", "Registro de proveedor")
    node(draw, (790, 180, 1065, 330), "Excel de proveedores", "Primera fuente operativa")
    node(draw, (1150, 180, 1530, 330), "Estado de resultados", "Segunda captura manual")
    arrow(draw, (345, 255), (430, 255))
    arrow(draw, (705, 255), (790, 255))
    arrow(draw, (1065, 255), (1150, 255))
    node(draw, (250, 470, 560, 620), "WANSOFT", "Ventas ya optimizadas")
    node(draw, (675, 470, 985, 620), "Reporte de ventas", "Exportación del periodo")
    arrow(draw, (560, 545), (675, 545))
    arrow(draw, (985, 545), (1310, 330))
    centered_text(draw, (1040, 480, 1525, 655), "Resultado: retraso, riesgo de error y poca trazabilidad", font(27, True), "#B33A3A")
    image.save(path, quality=95)


def save_demo_flow(path: Path):
    image, draw = base_canvas(1600, 900)
    draw.text((70, 45), "Demo disponible: automatización con revisión humana", font=font(42, True), fill="#17211C")
    boxes = [
        (55, 210, 300, 365, "Carga", "PDF, imagen o lote"),
        (365, 210, 610, 365, "Lectura", "Texto digital + OCR"),
        (675, 210, 920, 365, "Interpretación", "Campos y clasificación"),
        (985, 210, 1230, 365, "Revisión", "Capturista confirma"),
        (1295, 210, 1540, 365, "Publicación", "Registro aprobado"),
    ]
    for x1, y1, x2, y2, title, subtitle in boxes:
        node(draw, (x1, y1, x2, y2), title, subtitle, fill="#F8FBF9")
    for left, right in zip(boxes, boxes[1:]):
        arrow(draw, (left[2], 287), (right[0], 287))
    node(draw, (250, 550, 580, 770), "Corrección humana", "Proveedor, importes, categoría y tipo documental", fill="#FFF3DC", outline="#D6A928")
    node(draw, (690, 560, 1020, 730), "Analytics", "Aprobado y vista preliminar", fill="#E8F3ED")
    node(draw, (1130, 560, 1460, 730), "Excel", "CSV de consumo sin datos OCR", fill="#E8F3ED")
    arrow(draw, (1105, 365), (415, 560), color="#D6A928")
    arrow(draw, (1418, 365), (855, 560))
    arrow(draw, (1020, 645), (1130, 645))
    image.save(path, quality=95)


def save_target_flow(path: Path):
    image, draw = base_canvas(1600, 1020)
    draw.text((70, 40), "Arquitectura objetivo: una sola fuente financiera", font=font(42, True), fill="#17211C")
    sources = [
        (60, 170, 350, 300, "Facturas", "Proveedores y CEDIS"),
        (60, 370, 350, 500, "WANSOFT", "Ventas por restaurante"),
        (60, 570, 350, 700, "Órdenes de compra", "Asignación CEDIS"),
        (60, 770, 350, 900, "Otros ajustes", "Nómina y gastos sin factura"),
    ]
    for entry in sources:
        node(draw, entry[:4], entry[4], entry[5], fill="#F8FBF9")
    node(draw, (520, 255, 930, 475), "Plataforma financiera", "Validación, catálogos, auditoría y cierre semanal", fill="#E8F3ED")
    node(draw, (520, 605, 930, 825), "Modelo analítico", "Ventas, gastos, cuentas, áreas y comportamiento fijo/variable", fill="#FFF3DC", outline="#D6A928")
    node(draw, (1095, 160, 1535, 310), "Estado de resultados", "Por semana y restaurante", fill="#F8FBF9")
    node(draw, (1095, 410, 1535, 560), "Punto de equilibrio", "Ventas y tickets necesarios", fill="#F8FBF9")
    node(draw, (1095, 660, 1535, 810), "Dashboard ejecutivo", "Indicadores y alertas", fill="#F8FBF9")
    node(draw, (1095, 860, 1535, 970), "Excel exportable", "Formato de transición", fill="#F8FBF9")
    for _, y1, x2, y2, *_ in sources:
        arrow(draw, (x2, (y1 + y2) / 2), (520, 365))
    arrow(draw, (725, 475), (725, 605), color="#D6A928")
    for target_y in (235, 485, 735, 915):
        arrow(draw, (930, 715), (1095, target_y))
    image.save(path, quality=95)


def save_horizontal_chart(path: Path, title: str, labels: list[str], values: list[float], counts: list[int], note: str):
    width, height = 1500, 850
    image, draw = base_canvas(width, height)
    draw.text((70, 45), title, font=font(42, True), fill="#17211C")
    draw.text((70, 105), note, font=font(22), fill="#68746D")
    maximum = max(values) if values else 1
    top, bar_h, gap = 190, 70, 45
    label_x, chart_x, chart_width = 70, 490, 820
    for index, (label, value, count) in enumerate(zip(labels, values, counts)):
        y = top + index * (bar_h + gap)
        lines = wrap_lines(draw, label, font(25, True), 360)
        draw.multiline_text((label_x, y + 2), "\n".join(lines), font=font(25, True), fill="#17211C", spacing=5)
        draw.text((label_x, y + 43), f"{count} factura{'s' if count != 1 else ''}", font=font(19), fill="#68746D")
        draw.rounded_rectangle((chart_x, y + 9, chart_x + chart_width, y + bar_h - 9), radius=16, fill="#E8ECE9")
        fill_width = max(12, chart_width * value / maximum)
        color = "#176B45" if index == 0 else "#5C9875"
        draw.rounded_rectangle((chart_x, y + 9, chart_x + fill_width, y + bar_h - 9), radius=16, fill=color)
        draw.text((chart_x + chart_width + 25, y + 17), f"${value:,.0f}", font=font(24, True), fill="#17211C")
    image.save(path, quality=95)


def save_break_even_chart(path: Path):
    width, height = 1500, 850
    image, draw = base_canvas(width, height)
    draw.text((70, 45), "Punto de equilibrio: de la clasificación a una decisión operativa", font=font(40, True), fill="#17211C")
    draw.text((70, 103), "Esquema ilustrativo; los valores reales se calcularán por restaurante y semana", font=font(22), fill="#68746D")
    left, top, right, bottom = 150, 190, 1390, 720
    draw.line((left, top, left, bottom), fill="#68746D", width=4)
    draw.line((left, bottom, right, bottom), fill="#68746D", width=4)
    draw.text((47, 410), "Ingresos\ny costos", font=font(22, True), fill="#17211C", align="center")
    draw.text((620, 760), "Nivel de ventas", font=font(24, True), fill="#17211C")
    fixed_y = 570
    draw.line((left, fixed_y, right, fixed_y), fill="#D6A928", width=6)
    draw.text((right - 235, fixed_y - 42), "Costos fijos", font=font(22, True), fill="#A25D09")
    revenue_start, revenue_end = (left, bottom), (right, top + 30)
    cost_start, cost_end = (left, fixed_y), (right, top + 145)
    draw.line((revenue_start, revenue_end), fill="#176B45", width=8)
    draw.line((cost_start, cost_end), fill="#B33A3A", width=8)
    draw.text((right - 175, top + 5), "Ventas", font=font(24, True), fill="#176B45")
    draw.text((right - 255, top + 125), "Costos totales", font=font(24, True), fill="#B33A3A")
    # Intersection derived from the two line equations.
    x_ratio = (bottom - fixed_y) / ((bottom - revenue_end[1]) - (fixed_y - cost_end[1]))
    cross_x = left + x_ratio * (right - left)
    cross_y = bottom + x_ratio * (revenue_end[1] - bottom)
    draw.ellipse((cross_x - 14, cross_y - 14, cross_x + 14, cross_y + 14), fill="#D6A928", outline="#17211C", width=3)
    draw.line((cross_x, cross_y, cross_x, bottom), fill="#D6A928", width=4)
    draw.text((cross_x + 24, cross_y - 55), "Punto de equilibrio", font=font(25, True), fill="#17211C")
    draw.text((cross_x + 24, cross_y - 20), "Ventas = costos totales", font=font(20), fill="#68746D")
    draw.text((left + 35, top + 55), "Pérdida", font=font(28, True), fill="#B33A3A")
    draw.text((right - 330, top + 300), "Utilidad", font=font(28, True), fill="#176B45")
    image.save(path, quality=95)


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node_mar = tc_mar.find(qn(f"w:{margin}"))
        if node_mar is None:
            node_mar = OxmlElement(f"w:{margin}")
            tc_mar.append(node_mar)
        node_mar.set(qn("w:w"), str(value))
        node_mar.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: Iterable[int], indent=120):
    widths = list(widths)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run(run, size=None, color=INK, bold=None, italic=None, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def configure_styles(doc: Document):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for name, size, color, before, after in (
        ("Heading 1", 16, GREEN, 16, 8),
        ("Heading 2", 13, GREEN, 12, 6),
        ("Heading 3", 12, GREEN_DARK, 8, 4),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = rgb(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    caption = styles["Caption"]
    caption.font.name = "Calibri"
    caption.font.size = Pt(9)
    caption.font.color.rgb = rgb(MUTED)
    caption.font.italic = True
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.keep_with_next = True


def configure_numbering(doc: Document):
    numbering = doc.part.numbering_part.element
    for abstract_id, num_id, fmt, text_value in ((100, 100, "bullet", "\u2022"), (101, 101, "decimal", "%1.")):
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        level = OxmlElement("w:lvl")
        level.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        level.append(start)
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), fmt)
        level.append(num_fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), text_value)
        level.append(lvl_text)
        suffix = OxmlElement("w:suff")
        suffix.set(qn("w:val"), "space")
        level.append(suffix)
        level_justify = OxmlElement("w:lvlJc")
        level_justify.set(qn("w:val"), "left")
        level.append(level_justify)
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "720")
        tabs.append(tab)
        p_pr.append(tabs)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "720")
        ind.set(qn("w:hanging"), "360")
        p_pr.append(ind)
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "160")
        spacing.set(qn("w:line"), "280")
        spacing.set(qn("w:lineRule"), "auto")
        p_pr.append(spacing)
        level.append(p_pr)
        if fmt == "bullet":
            r_pr = OxmlElement("w:rPr")
            r_fonts = OxmlElement("w:rFonts")
            r_fonts.set(qn("w:ascii"), "Arial")
            r_fonts.set(qn("w:hAnsi"), "Arial")
            r_pr.append(r_fonts)
            level.append(r_pr)
        abstract.append(level)
        # OOXML requires every abstractNum before the concrete num elements.
        first_num = numbering.find(qn("w:num"))
        if first_num is None:
            numbering.append(abstract)
        else:
            numbering.insert(list(numbering).index(first_num), abstract)
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abstract_ref = OxmlElement("w:abstractNumId")
        abstract_ref.set(qn("w:val"), str(abstract_id))
        num.append(abstract_ref)
        numbering.append(num)


def add_list_item(doc: Document, text: str, ordered=False):
    paragraph = doc.add_paragraph()
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), "101" if ordered else "100")
    num_pr.append(ilvl)
    num_pr.append(num_id)
    p_pr.append(num_pr)
    paragraph.paragraph_format.left_indent = Inches(0.28)
    paragraph.paragraph_format.first_line_indent = Inches(-0.18)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.167
    set_run(paragraph.add_run(text), size=11)
    return paragraph


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(text, style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        set_run(p.add_run(bold_lead), bold=True)
        set_run(p.add_run(text[len(bold_lead):]))
    else:
        set_run(p.add_run(text))
    return p


def add_callout(doc, label: str, text: str, fill=GREEN_SOFT, accent=GREEN):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)
    set_cell_shading(cell, fill)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.keep_together = True
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    set_run(p.add_run(f"{label}: "), bold=True, color=accent)
    set_run(p.add_run(text), color=INK)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(1)


def add_table(doc, headers, rows, widths, aligns=None):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    set_repeat_table_header(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, LIGHT)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        if aligns:
            p.alignment = aligns[index]
        set_run(p.add_run(header), size=9.5, bold=True, color=GREEN_DARK)
    for row_data in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_data):
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[index].paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.08
            if aligns:
                p.alignment = aligns[index]
            set_run(p.add_run(str(value)), size=9.5)
    set_table_geometry(table, widths)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return table


def add_image(doc, path: Path, caption: str, width=6.3):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_together = True
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    shape = run.add_picture(str(path), width=Inches(width))
    doc_pr = shape._inline.docPr
    doc_pr.set("descr", caption)
    cap = doc.add_paragraph(caption, style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    relationship_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), GREEN)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.append(color)
    run_properties.append(underline)
    run.append(run_properties)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run(run, size=9, color=MUTED)


def configure_page(doc: Document):
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)
        section.different_first_page_header_footer = True
        header = section.header
        p = header.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        set_run(p.add_run("PROYECTO INTERNO  |  AUTOMATIZACIÓN FINANCIERA"), size=8.5, color=MUTED, bold=True)
        footer = section.footer
        table = footer.add_table(rows=1, cols=2, width=Inches(6.5))
        set_table_geometry(table, [7200, 2160], indent=0)
        table.cell(0, 0).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
        table.cell(0, 1).paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        set_run(table.cell(0, 0).paragraphs[0].add_run("Documento ejecutivo | Agosto 2026"), size=9, color=MUTED)
        page_p = table.cell(0, 1).paragraphs[0]
        set_run(page_p.add_run("Página "), size=9, color=MUTED)
        add_page_field(page_p)


def page_break(doc):
    doc.add_page_break()


def build_document():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    FIGURES_DIR.mkdir(parents=True, exist_ok=True)
    manual_flow = FIGURES_DIR / "flujo-actual.png"
    demo_flow = FIGURES_DIR / "flujo-demo.png"
    target_flow = FIGURES_DIR / "arquitectura-objetivo.png"
    category_chart = FIGURES_DIR / "gasto-categoria.png"
    provider_chart = FIGURES_DIR / "gasto-proveedor.png"
    break_even = FIGURES_DIR / "punto-equilibrio.png"
    save_manual_flow(manual_flow)
    save_demo_flow(demo_flow)
    save_target_flow(target_flow)
    save_horizontal_chart(
        category_chart,
        "Muestra preliminar: gasto por clasificación actual",
        ["Restaurante", "Variables", "Fijos", "Otros"],
        [350000.00, 12000.00, 8000.00, 3000.00],
        [26, 2, 1, 1],
        "Incluye aprobadas y pendientes; no representa todavía un cierre contable",
    )
    save_horizontal_chart(
        provider_chart,
        "Muestra preliminar: principales proveedores",
        ["Proveedor Ficticio Alfa", "Proveedor Ficticio Beta", "Proveedor Ficticio Gamma", "Proveedor Ficticio Delta", "Proveedor Ficticio Epsilon"],
        [150000.00, 120000.00, 20000.00, 15000.00, 10000.00],
        [9, 3, 4, 1, 2],
        "Importes sintéticos para validar visualización y revisión humana",
    )
    save_break_even_chart(break_even)

    doc = Document()
    configure_styles(doc)
    configure_numbering(doc)
    configure_page(doc)
    doc.core_properties.title = "Propuesta ejecutiva - Automatización financiera"
    doc.core_properties.subject = "Proyecto interno para cuatro restaurantes"
    doc.core_properties.author = "Equipo de proyecto"
    doc.core_properties.keywords = "facturación, OCR, estado de resultados, punto de equilibrio, WANSOFT"

    # Cover: proposal_centerpiece.
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(72)
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(14)
    set_run(kicker.add_run("PROYECTO INTERNO"), size=11, color=GOLD, bold=True)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    set_run(title.add_run("Automatización inteligente de facturación y análisis financiero"), size=27, color=GREEN_DARK, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(28)
    set_run(subtitle.add_run("De la captura manual al estado de resultados semanal y punto de equilibrio"), size=15, color=MUTED)
    cover_callout = doc.add_table(rows=1, cols=1)
    set_table_geometry(cover_callout, [9360])
    set_cell_shading(cover_callout.cell(0, 0), GREEN_SOFT)
    cp = cover_callout.cell(0, 0).paragraphs[0]
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_before = Pt(8)
    cp.paragraph_format.space_after = Pt(8)
    set_run(cp.add_run("Propuesta ejecutiva para dirección | Alcance inicial: cuatro restaurantes"), size=11.5, color=GREEN_DARK, bold=True)
    cover_space = doc.add_paragraph()
    cover_space.paragraph_format.space_after = Pt(75)
    prepared = doc.add_paragraph()
    prepared.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(prepared.add_run("Preparado para: Dirección"), size=11, color=INK, bold=True)
    date_p = doc.add_paragraph()
    date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(date_p.add_run("11 de agosto de 2026"), size=10.5, color=MUTED)
    confidentiality = doc.add_paragraph()
    confidentiality.alignment = WD_ALIGN_PARAGRAPH.CENTER
    confidentiality.paragraph_format.space_before = Pt(24)
    set_run(confidentiality.add_run("Documento de trabajo confidencial"), size=9.5, color=MUTED, italic=True)

    page_break(doc)
    add_heading(doc, "Resumen ejecutivo", 1)
    add_callout(
        doc,
        "Recomendación",
        "Adoptar la solución como proyecto interno para los cuatro restaurantes, conservando Excel como puente de transición y exportación, mientras la base de datos se convierte en la fuente oficial.",
    )
    add_body(doc, "Actualmente la información de proveedores se captura manualmente en Excel y después vuelve a transcribirse al estado de resultados semanal. El proceso funciona, pero duplica trabajo, retrasa el cierre y dificulta rastrear cómo se originó o corrigió cada cifra.")
    add_body(doc, "La demo construida ya demuestra el núcleo del cambio: recibir documentos, extraer información, separar tipos documentales, permitir corrección humana, aprobar registros y publicar únicamente datos de consumo en Analytics y Excel.")
    add_heading(doc, "Decisión solicitada", 2)
    add_body(doc, "Autorizar la evolución de la demo hacia un piloto interno controlado que conecte gastos aprobados, ventas provenientes de WANSOFT, ajustes financieros y reglas del estado de resultados semanal.")
    add_heading(doc, "Qué no se pretende hacer", 2)
    for item in (
        "No reemplazar WANSOFT ni volver a construir el punto de venta.",
        "No eliminar Excel de forma abrupta antes de conciliar resultados.",
        "No publicar información OCR o propuestas de IA como cifras financieras finales.",
        "No automatizar excepciones sin conservar una decisión humana auditable.",
    ):
        add_list_item(doc, item)
    add_heading(doc, "Resultado esperado", 2)
    add_body(doc, "Un cierre semanal más rápido y trazable, con estado de resultados por restaurante, capacidad de explicar cada gasto hasta su documento original y cálculo sistemático del punto de equilibrio.")

    page_break(doc)
    add_heading(doc, "1. Situación actual", 1)
    add_image(doc, manual_flow, "Figura 1. Flujo operativo actual: dos capturas manuales antes del análisis financiero.")
    add_heading(doc, "Problemas estructurales", 2)
    for item in (
        "Doble captura: factura a Excel de proveedores y Excel de proveedores a estado de resultados.",
        "Dependencia de conocimiento individual para clasificar y asignar cada gasto.",
        "Cambios de nombres, categorías y fórmulas difíciles de rastrear.",
        "Facturas del CEDIS sin identificación suficiente del restaurante receptor.",
        "El estado de resultados se actualiza después de que la operación ya ocurrió.",
    ):
        add_list_item(doc, item)
    add_callout(doc, "Principio de diseño", "La automatización debe eliminar transcripciones, no eliminar el control humano.", fill=GOLD_SOFT, accent="A25D09")

    page_break(doc)
    add_heading(doc, "2. Lo que existe hoy", 1)
    add_image(doc, demo_flow, "Figura 2. Flujo funcional de la demo disponible.")
    add_heading(doc, "Capacidades implementadas", 2)
    implemented = [
        ("Ingreso", "Carga individual, carga masiva y PDF con una factura por página."),
        ("Lectura", "Texto digital cuando existe y OCR local como respaldo para documentos escaneados."),
        ("Clasificación", "Separación entre factura fiscal, comprobantes de pago y otros documentos."),
        ("Revisión", "Corrección de proveedor, fecha, importes, categoría y subcategoría opcional."),
        ("Control", "El total es obligatorio; subtotal e impuestos pueden quedar pendientes sin falsearlos como cero."),
        ("Override", "El capturista puede corregir manualmente la decisión 'no genera gasto'."),
        ("Consumo", "Dashboard ejecutivo, vista aprobada/preliminar y descarga CSV compatible con Excel."),
    ]
    add_table(doc, ["Componente", "Capacidad actual"], implemented, [2100, 7260])
    add_body(doc, "El conjunto sintético de desarrollo valida documentos digitales y multipágina. Las pruebas automatizadas cubren el flujo principal, las correcciones y la publicación en Analytics.")

    page_break(doc)
    add_heading(doc, "3. Evidencia con datos sintéticos", 1)
    add_callout(doc, "Lectura correcta", "Las cifras de esta sección son preliminares: combinan registros aprobados y pendientes para demostrar capacidad analítica. No constituyen un estado financiero cerrado.", fill=GOLD_SOFT, accent="A25D09")
    add_image(doc, category_chart, "Gráfica 1. Distribución preliminar del gasto de la muestra por clasificación actual.", width=6.15)
    page_break(doc)
    add_image(doc, provider_chart, "Gráfica 2. Proveedores ficticios del conjunto sintético.", width=6.15)
    add_body(doc, "La visualización ya permite observar concentración de gasto y proveedores relevantes. El siguiente avance no es agregar más gráficas, sino convertir las clasificaciones transitorias en una estructura financiera consistente.")

    page_break(doc)
    add_heading(doc, "4. Modelo operativo recomendado", 1)
    add_image(doc, target_flow, "Figura 3. Arquitectura objetivo para el proyecto interno.", width=6.15)
    add_heading(doc, "Papel de cada herramienta", 2)
    roles = [
        ("WANSOFT", "Fuente de ventas", "Se conserva; se importan reportes por restaurante y periodo."),
        ("Plataforma", "Fuente operativa", "Recibe, revisa, clasifica y aprueba gastos."),
        ("Base de datos", "Fuente oficial", "Conserva hechos financieros, catálogos, reglas y auditoría."),
        ("Excel", "Puente y salida", "Se utiliza para conciliación, importaciones temporales y exportación."),
        ("Capturista", "Control humano", "Resuelve excepciones y confirma los datos antes de publicar."),
    ]
    add_table(doc, ["Elemento", "Función", "Decisión"], roles, [1700, 2200, 5460])
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    set_run(p.add_run("Referencia: "), size=9.5, color=MUTED, italic=True)
    add_hyperlink(p, "WANSOFT documenta reportes de resumen de ventas", "https://soporte.wansoft.net/hc/es-419/articles/18351925582861-Reportes-del-resumen-de-ventas")
    set_run(p.add_run(", que pueden funcionar como primera vía de integración."), size=9.5, color=MUTED, italic=True)

    page_break(doc)
    add_heading(doc, "5. Estructura financiera futura", 1)
    add_callout(doc, "Corrección necesaria", "Fijo/variable y barra/restaurante no deben vivir en el mismo catálogo. Son dimensiones distintas y se necesitan simultáneamente para analizar rentabilidad.")
    dimensions = [
        ("Cuenta del estado", "Renta, alimentos, nómina, mantenimiento", "¿Qué tipo de ingreso o gasto es?"),
        ("Comportamiento", "Fijo, variable, semivariable", "¿Cómo cambia cuando cambian las ventas?"),
        ("Área", "Restaurante, barra, administración, CEDIS", "¿Dónde se generó?"),
        ("Entidad", "Razón social y restaurante", "¿A quién corresponde?"),
        ("Periodo", "Semana financiera", "¿En qué cierre debe reconocerse?"),
    ]
    add_table(doc, ["Dimensión", "Ejemplos", "Pregunta que responde"], dimensions, [2000, 3300, 4060])
    add_heading(doc, "Estado de resultados semanal propuesto", 2)
    pnl_rows = [
        ("Ventas", "Ventas brutas - descuentos - cancelaciones", "WANSOFT"),
        ("Ventas netas", "Ingreso operativo del periodo", "Cálculo"),
        ("Costos variables", "Alimentos, bebidas, empaques y comisiones variables", "Facturas / ajustes"),
        ("Margen de contribución", "Ventas netas - costos variables", "Cálculo"),
        ("Gastos fijos y operativos", "Renta, nómina, servicios, mantenimiento y administración", "Facturas / nómina"),
        ("Utilidad operativa", "Margen de contribución - gastos fijos y operativos", "Cálculo"),
        ("Otros ingresos y gastos", "Conceptos no operativos definidos por dirección", "Ajustes"),
        ("Resultado semanal", "Resultado final por restaurante", "Cálculo"),
    ]
    add_table(doc, ["Renglón", "Definición", "Fuente"], pnl_rows, [2200, 5000, 2160])
    add_body(doc, "El estado de resultados y el flujo de efectivo deben mantenerse separados. La fecha de emisión, la semana de consumo y la fecha de pago pueden ser diferentes.")

    page_break(doc)
    add_heading(doc, "6. Punto de equilibrio", 1)
    add_image(doc, break_even, "Gráfica 3. Relación conceptual entre ventas, costos y punto de equilibrio.", width=6.2)
    add_heading(doc, "Cálculo base", 2)
    formula_rows = [
        ("Margen de contribución", "Ventas netas - costos variables"),
        ("Porcentaje de contribución", "Margen de contribución / ventas netas"),
        ("Ventas de equilibrio", "Costos fijos / porcentaje de contribución"),
        ("Tickets de equilibrio", "Ventas de equilibrio / ticket promedio WANSOFT"),
    ]
    add_table(doc, ["Indicador", "Fórmula"], formula_rows, [3300, 6060])
    add_body(doc, "El valor ejecutivo del cálculo no es solamente conocer una cifra. Permitirá mostrar por restaurante cuánto falta para cubrir costos, cuántos tickets adicionales se requieren y cómo cambia el resultado cuando varía el costo de insumos.")
    add_callout(doc, "Condición", "El punto de equilibrio será confiable únicamente cuando los gastos estén separados por comportamiento fijo, variable o semivariable.", fill=GOLD_SOFT, accent="A25D09")

    page_break(doc)
    add_heading(doc, "7. Roadmap funcional", 1)
    add_body(doc, "La implementación debe avanzar por entregables verificables, sin eliminar el proceso actual antes de comprobar que el nuevo resultado coincide.")
    roadmap = [
        ("1. Consolidar la demo", "Cuatro razones sociales, catálogos definitivos, permisos, respaldos y reglas de revisión.", "Gasto aprobado y auditable."),
        ("2. Integrar finanzas", "Importación WANSOFT, estructura exacta del estado de resultados, nómina y ajustes.", "Borrador semanal automático."),
        ("3. Resolver CEDIS", "Órdenes de compra, sugerencia de restaurante y confirmación humana.", "Menos gastos sin asignar."),
        ("4. Cerrar en paralelo", "Comparar proceso manual contra plataforma y convertir diferencias en reglas.", "Cierre conciliado y explicable."),
        ("5. Operación interna", "Dashboard, punto de equilibrio, exportación y gobierno de cierre.", "Fuente financiera institucional."),
    ]
    add_table(doc, ["Etapa", "Contenido", "Resultado"], roadmap, [1900, 4740, 2720])
    add_heading(doc, "Indicadores para evaluar el piloto", 2)
    for item in (
        "Porcentaje de documentos que llegan a revisión con proveedor, fecha y total identificados.",
        "Tiempo manual desde recepción hasta publicación.",
        "Facturas pendientes o sin restaurante al momento del cierre.",
        "Número y causa de ajustes entre el estado manual y el generado.",
        "Capacidad de explicar cada renglón hasta el documento de origen.",
    ):
        add_list_item(doc, item)

    page_break(doc)
    add_heading(doc, "8. Gobierno, riesgos y siguientes insumos", 1)
    add_heading(doc, "Controles indispensables", 2)
    for item in (
        "Cada semana debe tener estado borrador, revisión y cierre aprobado.",
        "Toda corrección debe conservar usuario, dato anterior, dato nuevo y motivo.",
        "Los catálogos y reglas financieras deben tener un responsable de negocio.",
        "Analytics y Excel deben consumir únicamente registros aprobados; la vista preliminar debe permanecer identificada.",
        "Un cierre reabierto debe conservar versiones y explicación del cambio.",
    ):
        add_list_item(doc, item)
    add_heading(doc, "Riesgos que requieren atención", 2)
    risks = [
        ("Asignación CEDIS", "La factura puede no identificar al receptor.", "Cruzar órdenes de compra y solicitar confirmación humana."),
        ("Semana financiera", "Emisión, consumo y pago pueden pertenecer a semanas distintas.", "Definir una regla administrativa única."),
        ("Integración WANSOFT", "El mecanismo disponible puede ser reporte y no API.", "Iniciar con exportación controlada y automatizar después."),
        ("Catálogos", "Una clasificación inconsistente afecta utilidad y equilibrio.", "Separar cuenta, área y comportamiento."),
        ("Cambio operativo", "Excel es conocido y flexible.", "Ejecutar cierres paralelos antes de retirarlo como fuente."),
    ]
    add_table(doc, ["Riesgo", "Impacto", "Respuesta"], risks, [1900, 3100, 4360])
    add_heading(doc, "Información necesaria para continuar", 2)
    for item in (
        "Plantilla vigente del estado de resultados semanal.",
        "Exportación representativa de ventas desde WANSOFT.",
        "Las tres razones sociales restantes y sus alias.",
        "Excel de órdenes de compra del CEDIS.",
        "Regla actual para decidir la semana a la que pertenece cada gasto.",
        "Fuentes de nómina y otros conceptos no respaldados por factura.",
    ):
        add_list_item(doc, item)
    add_callout(doc, "Cierre ejecutivo", "La demo ya valida la viabilidad técnica. La siguiente decisión es convertirla en un sistema interno de cierre financiero, con Excel como transición y WANSOFT como fuente de ventas.")

    doc.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    build_document()
